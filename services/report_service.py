# -*- coding: utf-8 -*-
"""报告导出服务：把问答/对比结果导出为 Word(docx) 或 Markdown。

对外方法：
    build_markdown(payload) -> str
    build_docx(payload)     -> bytes
    build_session_docx(history, session_id, title) -> bytes
"""
from __future__ import annotations

import io
import re
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

DISCLAIMER = ("本报告由 FinSight 酒研助手基于公开年报数据自动生成，"
              "仅供研究学习参考，不构成任何投资建议。")


def _clean_inline(text: str) -> str:
    """去掉 Markdown 行内标记（**加粗**、`代码`、[文字](链接)）。"""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    return text.strip()


def _is_table_sep(line: str) -> bool:
    return "-" in line and re.fullmatch(r"\|?[\s:|-]+\|?", line.strip()) is not None


def _add_markdown_block(doc, markdown_text: str) -> None:
    """把简化 Markdown 写进 docx（标题/列表/引用/表格/普通段落）。"""
    lines = (markdown_text or "").splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        # 表格：连续的 | 行
        if line.startswith("|") and line.endswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not _is_table_sep(lines[i]):
                    block.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            if block:
                cols = max(len(r) for r in block)
                table = doc.add_table(rows=len(block), cols=cols)
                table.style = "Light Grid Accent 1"
                for r, cells in enumerate(block):
                    for c in range(cols):
                        table.cell(r, c).text = cells[c] if c < len(cells) else ""
            continue
        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = min(len(m.group(1)) + 1, 4)
            doc.add_heading(_clean_inline(m.group(2)), level=level)
            i += 1
            continue
        # 引用
        if line.startswith(">"):
            p = doc.add_paragraph()
            run = p.add_run(_clean_inline(line.lstrip("> ")))
            run.italic = True
            run.font.color.rgb = RGBColor(0x8A, 0x6D, 0x28)
            i += 1
            continue
        # 列表
        m = re.match(r"^[-*+·]\s+(.*)$", line)
        if m:
            doc.add_paragraph(_clean_inline(m.group(1)), style="List Bullet")
            i += 1
            continue
        m = re.match(r"^(\d+)[.、]\s*(.*)$", line)
        if m:
            doc.add_paragraph(_clean_inline(m.group(2)), style="List Number")
            i += 1
            continue
        p = doc.add_paragraph(_clean_inline(line))
        for run in p.runs:
            run.font.size = Pt(10.5)
        i += 1


def build_markdown(payload: dict) -> str:
    title = payload.get("title") or "FinSight 酒研助手 · 问答报告"
    created = payload.get("created_at") or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    parts = [f"# {title}", "", f"> 生成时间：{created}", ""]
    question = payload.get("question")
    answer = payload.get("answer")
    if question:
        parts += ["## 问题", "", question, "", "## 回答", "", answer or "", ""]
    for item in payload.get("history") or []:
        parts += [f"## 问题：{item.get('question', '')}", "", item.get("answer", ""), ""]
    parts += ["---", "", DISCLAIMER, ""]
    return "\n".join(parts)


def _new_document(title: str, subtitle: str) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10.5)
    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(subtitle)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in sub.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    return doc


def build_docx(payload: dict) -> bytes:
    title = payload.get("title") or "FinSight 酒研助手 · 问答报告"
    created = payload.get("created_at") or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    mode = payload.get("mode") or "auto"
    doc = _new_document(title, f"模式：{mode} ｜ 生成时间：{created}")

    question = payload.get("question")
    answer = payload.get("answer")
    if question:
        doc.add_heading("问题", level=1)
        doc.add_paragraph(question)
        doc.add_heading("回答", level=1)
        _add_markdown_block(doc, answer or "")
    for idx, item in enumerate(payload.get("history") or [], 1):
        doc.add_heading(f"第 {idx} 轮 · 问题", level=1)
        doc.add_paragraph(item.get("question", ""))
        doc.add_heading("回答", level=2)
        _add_markdown_block(doc, item.get("answer", ""))

    doc.add_paragraph("")
    note = doc.add_paragraph(DISCLAIMER)
    for run in note.runs:
        run.italic = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x9A, 0xA5, 0xB4)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_session_docx(history: list[dict], session_id: str = "", title: str | None = None) -> bytes:
    payload = {
        "title": title or "FinSight 酒研助手 · 会话报告",
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "history": history,
    }
    return build_docx(payload)
